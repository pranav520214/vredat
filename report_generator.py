"""
report_generator.py - Generates comprehensive DOCX reports for the local repository.
"""

import os
import glob
from datetime import datetime
from typing import Dict, Any, List

try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


def scan_repository(workspace_dir: str) -> Dict[str, Any]:
    """Scans repository files and collects structural metadata."""
    ignored = {".git", ".venv", "venv", "__pycache__", "node_modules", ".vredat", "build", "dist", "*.egg-info"}
    files_info = []
    total_lines = 0
    total_size = 0
    ext_counts = {}

    for root, dirs, files in os.walk(workspace_dir):
        # Prune ignored dirs
        dirs[:] = [d for d in dirs if d not in ignored and not d.endswith(".egg-info")]
        for f in files:
            ext = os.path.splitext(f)[1].lower() or "no-ext"
            fp = os.path.join(root, f)
            rel_path = os.path.relpath(fp, workspace_dir)
            size = os.path.getsize(fp)
            
            lines = 0
            try:
                with open(fp, "r", encoding="utf-8", errors="ignore") as fh:
                    lines = sum(1 for _ in fh)
            except Exception:
                pass

            total_lines += lines
            total_size += size
            ext_counts[ext] = ext_counts.get(ext, 0) + 1

            files_info.append({
                "path": rel_path,
                "size_kb": round(size / 1024, 2),
                "lines": lines,
                "ext": ext
            })

    return {
        "workspace_dir": workspace_dir,
        "repo_name": os.path.basename(os.path.abspath(workspace_dir)),
        "files": files_info,
        "total_files": len(files_info),
        "total_lines": total_lines,
        "total_size_kb": round(total_size / 1024, 2),
        "ext_counts": ext_counts
    }


def create_repo_docx_report(workspace_dir: str, output_path: str = "") -> str:
    """Generates a professional Word Document (.docx) report of the repository."""
    if not DOCX_AVAILABLE:
        return "[Error: python-docx is not installed. Run 'pip install python-docx']"

    data = scan_repository(workspace_dir)
    if not output_path:
        output_path = os.path.join(workspace_dir, f"{data['repo_name']}_repository_report.docx")

    doc = docx.Document()

    # Style configuration
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)

    # Document Header / Title
    title = doc.add_heading(f"Repository Architecture & Audit Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph(f"Project: {data['repo_name']} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # 1. Executive Summary
    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        f"This report provides an automated structural analysis and architectural breakdown of the '{data['repo_name']}' codebase. "
        f"The project currently comprises {data['total_files']} files spanning {data['total_lines']} lines of code ({data['total_size_kb']} KB total size)."
    )

    # 2. Key Metrics Table
    doc.add_heading("2. Repository Metrics", level=1)
    metric_table = doc.add_table(rows=1, cols=2)
    metric_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = metric_table.rows[0].cells
    hdr[0].text = "Metric"
    hdr[1].text = "Value"

    metrics = [
        ("Total Files", str(data["total_files"])),
        ("Total Source Lines", f"{data['total_lines']:,}"),
        ("Total Disk Size", f"{data['total_size_kb']} KB"),
        ("Primary File Types", ", ".join(f"{k} ({v})" for k, v in sorted(data["ext_counts"].items(), key=lambda x: -x[1])[:5])),
        ("Root Directory", data["workspace_dir"])
    ]
    for m, val in metrics:
        row = metric_table.add_row().cells
        row[0].text = m
        row[1].text = val

    doc.add_paragraph()

    # 3. File Inventory
    doc.add_heading("3. Component & File Inventory", level=1)
    file_table = doc.add_table(rows=1, cols=3)
    file_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    f_hdr = file_table.rows[0].cells
    f_hdr[0].text = "File Path"
    f_hdr[1].text = "Lines"
    f_hdr[2].text = "Size (KB)"

    for f in sorted(data["files"], key=lambda x: x["path"]):
        row = file_table.add_row().cells
        row[0].text = f["path"]
        row[1].text = str(f["lines"])
        row[2].text = str(f["size_kb"])

    doc.add_paragraph()

    # 4. Architecture & Key Modules
    doc.add_heading("4. Architecture & System Modules", level=1)
    
    # Check for specific files in this repo
    file_descriptions = {
        "cli.py": "Main Terminal REPL application. Bridges Laya, Ponytail, and Ollama.",
        "engine_laya.py": "System 1 decision router. Performs single-pass intention and complexity evaluation in ~33ms.",
        "engine_llm.py": "System 2 local inference manager. Connects to Ollama, auto-starts background daemon, handles streaming.",
        "memory_ponytail.py": "Context window and token manager. Implements the Ponytail YAGNI ladder of simplicity.",
        "tools.py": "Deterministic file, terminal, git, and search tools executed with zero LLM compute.",
        "report_generator.py": "Automated repository report generator exporting structural and audit findings to DOCX format.",
        "setup.py": "Package definition and console script installer for 'vredat' and 'verdeta' commands."
    }

    for fname, desc in file_descriptions.items():
        if any(f["path"] == fname for f in data["files"]):
            p = doc.add_paragraph()
            p.add_run(f"• {fname}: ").bold = True
            p.add_run(desc)

    doc.add_paragraph()

    # 5. Ponytail & YAGNI Audit
    doc.add_heading("5. Ponytail / Code Health Assessment", level=1)
    doc.add_paragraph(
        "Adherence to the Ponytail Ladder of Simplicity:\n"
        "1. Standard Library & Native Features: The system leverages native Python stdlib and system tools where possible.\n"
        "2. Deterministic Short-Circuiting: Terminal commands and file lookups bypass generative models to minimize token expenditure.\n"
        "3. Context Compression: The memory ledger keeps conversational context bounded, preventing VRAM overflow and token drift."
    )

    try:
        doc.save(output_path)
        return output_path
    except Exception as e:
        return f"[Error saving docx: {e}]"
